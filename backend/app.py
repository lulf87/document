import os
from fastapi import FastAPI, File, UploadFile, Form, Query, Body
from fastapi.responses import JSONResponse, FileResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from typing import List
import shutil
import tempfile
import docx
import openpyxl
import pdfplumber
from pdf2image import convert_from_path
import pytesseract
from PIL import Image
from datetime import datetime
import json
import requests
from sentence_transformers import SentenceTransformer
import chromadb
from chromadb.config import Settings
from chromadb.utils import embedding_functions

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(PROJECT_ROOT, "uploaded_files")
CHROMA_DIR = os.path.join(UPLOAD_DIR, "chroma_db")

print("正在初始化目录...")
# 创建上传文件目录
os.makedirs(UPLOAD_DIR, exist_ok=True)
# 创建ChromaDB持久化目录
os.makedirs(CHROMA_DIR, exist_ok=True)
print(f"目录初始化完成：\n- 上传目录：{UPLOAD_DIR}\n- ChromaDB目录：{CHROMA_DIR}")

METADATA_FILE = os.path.join(UPLOAD_DIR, 'metadata.json')
DEEPSEEK_KEY_FILE = os.path.join(UPLOAD_DIR, 'deepseek.key')

# 初始化ChromaDB
print("ChromaDB持久化路径：", CHROMA_DIR)
chroma_client = chromadb.PersistentClient(path=CHROMA_DIR)
COLLECTION_NAME = 'doc_chunks'
if COLLECTION_NAME not in [c.name for c in chroma_client.list_collections()]:
    print(f"创建新的ChromaDB collection: {COLLECTION_NAME}")
    chroma_client.create_collection(COLLECTION_NAME)
collection = chroma_client.get_collection(COLLECTION_NAME)

# 输出已加载文档块数量
print(f"ChromaDB collection '{COLLECTION_NAME}' 已加载文档块数量：{collection.count()}")

# 初始化embedding模型
embed_model = SentenceTransformer('../paraphrase-multilingual-MiniLM-L12-v2')

# 分块参数
CHUNK_SIZE = 3000
CHUNK_OVERLAP = 300

def load_metadata():
    if not os.path.exists(METADATA_FILE):
        return {}
    with open(METADATA_FILE, 'r', encoding='utf-8') as f:
        return json.load(f)

def save_metadata(meta):
    with open(METADATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

def save_deepseek_key(key):
    with open(DEEPSEEK_KEY_FILE, 'w') as f:
        f.write(key.strip())

def load_deepseek_key():
    if not os.path.exists(DEEPSEEK_KEY_FILE):
        return None
    with open(DEEPSEEK_KEY_FILE, 'r') as f:
        return f.read().strip()

# 解析Word文档
def parse_docx(file_path):
    doc = docx.Document(file_path)
    return '\n'.join([para.text for para in doc.paragraphs])

# 解析Excel文档
def parse_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    text = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text.append('\t'.join([str(cell) if cell is not None else '' for cell in row]))
    return '\n'.join(text)

# 解析PDF文档（含OCR）
def parse_pdf(file_path):
    text = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)
            else:
                # OCR
                images = convert_from_path(file_path, first_page=page.page_number, last_page=page.page_number)
                for image in images:
                    ocr_text = pytesseract.image_to_string(image, lang='chi_sim+eng')
                    text.append(ocr_text)
    return '\n'.join(text)

# 文件解析主入口
def extract_text(file_path, ext):
    if ext == '.docx':
        return parse_docx(file_path)
    elif ext == '.xlsx':
        return parse_xlsx(file_path)
    elif ext == '.pdf':
        return parse_pdf(file_path)
    else:
        return ''

# 文本分块
def split_text(text, chunk_size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks

# 上传文档时自动分块、embedding入库
@app.post("/upload/")
async def upload_file(file: UploadFile = File(...)):
    print(f"\n开始处理文件上传：{file.filename}")
    
    # 1. 检查文件类型
    ext = os.path.splitext(file.filename)[-1].lower()
    if ext not in ['.docx', '.xlsx', '.pdf']:
        return JSONResponse(status_code=400, content={"error": "仅支持docx、xlsx、pdf文件"})
    
    # 2. 使用原始文件名，但添加防重名机制
    original_filename = file.filename
    save_filename = original_filename
    base, ext = os.path.splitext(original_filename)
    counter = 1
    
    # 如果文件名已存在，自动加上序号
    while os.path.exists(os.path.join(UPLOAD_DIR, save_filename)):
        save_filename = f"{base}_{counter}{ext}"
        counter += 1
    
    print(f"保存文件名：{save_filename}")
    
    # 3. 保存文件
    file_path = os.path.join(UPLOAD_DIR, save_filename)
    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    
    try:
        # 4. 解析文本
        print("开始解析文本...")
        text = extract_text(file_path, ext)
        print(f"文本解析完成，长度：{len(text)} 字符")
        
        # 5. RAG: 分块并入库
        print("开始文本分块...")
        chunks = split_text(text)
        print(f"分块完成，共 {len(chunks)} 个块")
        
        chunk_ids = []
        metadatas = []
        for i, chunk in enumerate(chunks):
            chunk_id = f"{save_filename}_{i}"
            chunk_ids.append(chunk_id)
            metadatas.append({
                "filename": save_filename,
                "original_filename": original_filename,
                "chunk_index": i,
                "group": "未分组"
            })
        
        if chunks:
            print("开始生成embedding...")
            embeddings = embed_model.encode(chunks).tolist()
            print(f"embedding生成完成，开始写入ChromaDB...")
            
            print("metadatas:", metadatas)
            
            collection.upsert(
                ids=chunk_ids,
                embeddings=embeddings,
                documents=chunks,
                metadatas=metadatas
            )
            print(f"ChromaDB写入完成")
            print(f"当前ChromaDB中总文档块数：{collection.count()}")
        
        # 6. 保存原始文件名到metadata
        meta = load_metadata()
        meta[save_filename] = {
            "original_filename": original_filename,
            "group": "未分组"
        }
        save_metadata(meta)
        
        test = collection.get(where={"filename": save_filename})
        print("实际存储的metadata:", test.get("metadatas", []))
        
        print(f"文件处理完成：{save_filename}\n")
        
        return {
            "filename": save_filename,
            "original_filename": original_filename,
            "text": text
        }
        
    except Exception as e:
        print(f"处理文件时出错：{str(e)}")
        os.remove(file_path)
        return JSONResponse(status_code=500, content={"error": f"文件解析失败: {str(e)}"})

@app.post("/set_deepseek_key/")
def set_deepseek_key(key: str = Form(...)):
    save_deepseek_key(key)
    return {"success": True}

@app.post("/test_deepseek_key/")
def test_deepseek_key(key: str = Form(...)):
    # 用一个简单的prompt测试key有效性
    url = "https://api.deepseek.com/v1/chat/completions"
    headers = {"Authorization": f"Bearer {key.strip()}", "Content-Type": "application/json"}
    data = {
        "model": "deepseek-chat",
        "messages": [{"role": "user", "content": "你好"}],
        "max_tokens": 20
    }
    try:
        resp = requests.post(url, headers=headers, json=data, timeout=10)
        if resp.status_code == 200:
            return {"success": True}
        else:
            return JSONResponse(status_code=400, content={"error": f"API Key无效: {resp.text}"})
    except Exception as e:
        return JSONResponse(status_code=500, content={"error": f"连接失败: {str(e)}"})

# RAG问答接口
@app.post("/ask/")
async def ask_ai(body: dict = Body(...)):
    print("\n开始处理问答请求...")
    doc_count = collection.count()
    print(f"当前ChromaDB中的文档块数量：{doc_count}")
    if doc_count == 0:
        return JSONResponse(status_code=400, content={"error": "当前没有可用的文档，请先上传文档。"})
    messages = body.get("messages", [])
    if not messages:
        return JSONResponse(status_code=400, content={"error": "消息不能为空"})

    # 获取最新的用户问题
    user_question = messages[-1]["content"]
    print(f"用户问题：{user_question}")

    # 先生成query_embedding，后续所有检索都用
    query_embedding = embed_model.encode(user_question).tolist()

    # 新增：支持group参数
    group = body.get("group")
    if group:
        print(f"仅检索分组：{group}")
        where = {"group": group}
        query_kwargs = dict(query_embeddings=[query_embedding], n_results=100, where=where)
        results = collection.query(**query_kwargs)
    else:
        query_kwargs = dict(query_embeddings=[query_embedding], n_results=100)
        results = collection.query(**query_kwargs)
    
    if not results['documents'][0]:
        print("未找到相关文档内容")
        return JSONResponse(status_code=400, content={"error": "未找到相关文档内容"})
    
    # 获取所有相关的文档内容及其来源文件
    relevant_docs = results['documents'][0]  # 所有相关文档块
    print(f"找到 {len(relevant_docs)} 个相关文档块")
    
    source_files = []  # 用于存储不重复的来源文件
    source_files_set = set()  # 用于去重
    
    # 收集所有不重复的来源文件
    for metadata in results['metadatas'][0]:
        if metadata['filename'] not in source_files_set:
            source_files_set.add(metadata['filename'])
            source_files.append(metadata['filename'])
    
    print(f"相关文档来源：{', '.join(source_files)}")
    
    # 合并所有相关文档块的内容，并限制最大字符数
    combined_docs = "\n\n---\n\n".join(relevant_docs)
    MAX_DOC_CHARS = 50000
    if len(combined_docs) > MAX_DOC_CHARS:
        print(f"相关文档内容过长，已截断至{MAX_DOC_CHARS}字符")
        combined_docs = combined_docs[:MAX_DOC_CHARS] + "\n...（内容已截断）..."
    
    # 2. 构建system prompt
    system_prompt = f"""你是一个专业的文档问答助手。请基于以下文档内容，回答用户的问题。
如果问题无法从文档内容中得到答案，请明确告知。
请保持专业、准确、简洁的回答风格。

相关文档内容：
{combined_docs}

其他注意事项：
1. 只基于提供的文档内容回答
2. 如果文档内容不足以回答问题，请直接说明
3. 保持客观、准确，不要添加个人观点
4. 如果发现多个文档中的信息有冲突，请指出这些冲突
"""
    
    # 3. 调用DeepSeek API
    key = load_deepseek_key()
    if not key:
        return JSONResponse(status_code=400, content={"error": "请先设置DeepSeek API Key"})
    
    print(f"使用的API Key: {key[:6]}...{key[-4:]}")  # 只打印密钥的前6位和后4位
    
    url = "https://api.deepseek.com/v1/chat/completions"
    headers = {"Authorization": f"Bearer {key.strip()}", "Content-Type": "application/json"}
    
    # 打印请求数据（不包含敏感信息）
    request_data = {
        "model": "deepseek-chat",
        "messages": [
            {"role": "system", "content": "系统提示信息..."},  # 不打印实际内容
            {"role": "user", "content": user_question}
        ]
    }
    print(f"请求数据: {json.dumps(request_data, ensure_ascii=False, indent=2)}")
    
    data = {
        "model": "deepseek-chat",
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_question}
        ],
        "temperature": 0.7,
        "stream": False
    }
    
    try:
        print(f"开始调用DeepSeek API...")
        response = requests.post(url, headers=headers, json=data)
        print(f"API响应状态码: {response.status_code}")
        print(f"API响应内容: {response.text}")
        
        if response.status_code != 200:
            error_msg = f"DeepSeek API调用失败: HTTP {response.status_code}"
            try:
                error_detail = response.json()
                error_msg += f"\n详细信息: {json.dumps(error_detail, ensure_ascii=False, indent=2)}"
            except:
                error_msg += f"\n响应内容: {response.text}"
            print(error_msg)
            return JSONResponse(status_code=response.status_code, content={"error": error_msg})
            
        response.raise_for_status()
        answer = response.json()["choices"][0]["message"]["content"]
        
        # 返回答案和所有来源文件信息
        return {
            "answer": answer,
            "source_files": source_files  # 返回所有相关文件名列表
        }
        
    except requests.exceptions.RequestException as e:
        error_msg = f"网络请求异常: {str(e)}"
        print(error_msg)
        return JSONResponse(status_code=500, content={"error": error_msg})
    except Exception as e:
        error_msg = f"处理异常: {str(e)}"
        print(error_msg)
        return JSONResponse(status_code=500, content={"error": error_msg})

@app.get("/files/")
def list_files():
    meta = load_metadata()
    files = []
    for fname in os.listdir(UPLOAD_DIR):
        if fname == 'metadata.json' or fname == 'deepseek.key':
            continue
        fpath = os.path.join(UPLOAD_DIR, fname)
        if os.path.isfile(fpath):
            stat = os.stat(fpath)
            file_meta = meta.get(fname, {})
            files.append({
                "filename": fname,
                "original_filename": file_meta.get('original_filename', fname),
                "size": stat.st_size,
                "upload_time": datetime.fromtimestamp(stat.st_ctime).strftime('%Y-%m-%d %H:%M:%S'),
                "group": file_meta.get('group', '未分组')
            })
    files.sort(key=lambda x: x["upload_time"], reverse=True)
    return {"files": files}

@app.post("/delete/")
def delete_file(filename: str = Query(...)):
    fpath = os.path.join(UPLOAD_DIR, filename)
    if os.path.exists(fpath) and os.path.isfile(fpath):
        os.remove(fpath)
        # 同步删除ChromaDB中所有属于该文件的分块
        print(f"同步删除ChromaDB中属于{filename}的所有分块...")
        results = collection.get(where={"filename": filename})
        ids = results.get("ids", [])
        if ids:
            collection.delete(ids=ids)
            print(f"已删除{len(ids)}个分块")
        else:
            print("未找到相关分块，无需删除")
        return {"success": True}
    else:
        return JSONResponse(status_code=404, content={"error": "文件不存在"})

@app.post("/rename/")
def rename_file(old_name: str = Query(...), new_name: str = Query(...)):
    old_path = os.path.join(UPLOAD_DIR, old_name)
    new_path = os.path.join(UPLOAD_DIR, new_name)
    if not os.path.exists(old_path):
        return JSONResponse(status_code=404, content={"error": "原文件不存在"})
    if os.path.exists(new_path):
        return JSONResponse(status_code=400, content={"error": "新文件名已存在"})
    os.rename(old_path, new_path)
    meta = load_metadata()
    if old_name in meta:
        meta[new_name] = meta.pop(old_name)
    save_metadata(meta)
    return {"success": True}

@app.post("/group/")
def set_group(filename: str = Query(...), group: str = Query(...)):
    fpath = os.path.join(UPLOAD_DIR, filename)
    if not os.path.exists(fpath):
        return JSONResponse(status_code=404, content={"error": "文件不存在"})
    meta = load_metadata()
    if filename not in meta:
        meta[filename] = {}
    meta[filename]['group'] = group
    save_metadata(meta)

    # 同步更新ChromaDB中所有相关块的group字段
    print(f"同步更新ChromaDB中{filename}的所有块的分组为：{group}")
    # 查询所有属于该文件的块
    results = collection.get(where={"filename": filename})
    ids = results.get("ids", [])
    if ids:
        # 批量更新group字段
        collection.update(
            ids=ids,
            metadatas=[{"group": group} for _ in ids]
        )
        print(f"已更新{len(ids)}个块的分组信息")
    else:
        print(f"未找到属于{filename}的块，无需更新ChromaDB分组")
    return {"success": True}

@app.get("/groups/")
def get_groups():
    meta = load_metadata()
    groups = {}
    for fname, info in meta.items():
        group = info.get('group', '未分组')
        if group not in groups:
            groups[group] = []
        groups[group].append(fname)
    return {"groups": groups}

@app.post("/fix_group_metadata/")
def fix_group_metadata():
    print("开始修复所有块的group字段...")
    # 获取所有块
    all_chunks = collection.get()
    ids = all_chunks.get("ids", [])
    metadatas = all_chunks.get("metadatas", [])
    update_ids = []
    for i, meta in enumerate(metadatas):
        group = meta.get("group", None)
        if group is None or group == "" or group == []:
            update_ids.append(ids[i])
    if update_ids:
        collection.update(
            ids=update_ids,
            metadatas=[{"group": "未分组"} for _ in update_ids]
        )
        print(f"已修复{len(update_ids)}个块的group字段为空的问题")
        return {"success": True, "fixed": len(update_ids)}
    else:
        print("所有块的group字段均已正确，无需修复")
        return {"success": True, "fixed": 0}

# 挂载前端静态文件
frontend_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "frontend", "dist")
if os.path.exists(frontend_path):
    app.mount("/", StaticFiles(directory=frontend_path, html=True), name="frontend")

@app.get("/")
async def read_root():
    return FileResponse(os.path.join(frontend_path, "index.html"))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8004) 